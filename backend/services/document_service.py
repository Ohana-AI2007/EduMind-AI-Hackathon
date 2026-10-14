"""Upload validation, extraction (PDF/PPTX/TXT/MD/video), chunking and indexing."""
import logging
import os
import re
import shutil
import unicodedata
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from sqlalchemy import select
from sqlalchemy.orm import Session

from backend import llm
from backend.config import settings
from backend.database import SessionLocal
from backend.errors import (AIConfigError, AIServiceError, BadRequestError, ExtractionError, PayloadTooLargeError)
from backend.models import ActivityLog, Chunk, Document, new_id
from backend.services import embedding_service, video_service

log = logging.getLogger(__name__)
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v", ".mp3", ".wav", ".m4a"}
SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md"} | VIDEO_EXT


@dataclass
class Segment:
    text: str
    topic: str
    page: Optional[int] = None
    slide: Optional[int] = None
    start_sec: Optional[float] = None
    end_sec: Optional[float] = None
    content_type: str = "text"
    presplit: bool = False


# ---------- helpers ----------
def safe_filename(name: str) -> str:
    name = os.path.basename((name or "").replace("\\", "/"))
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    name = re.sub(r"\s+", " ", name)
    return name[:120] or "upload"


def chunk_text(text: str, size: int, overlap: int) -> list:
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            window = text[start:end]
            cut = max(window.rfind("\n\n"), window.rfind(". "), window.rfind("\n"))
            if cut > size * 0.5:
                end = start + cut + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


_HEAD_KW = re.compile(r"^(chapter|section|unit|lesson|topic|module)\s+[\w.]+\b.*$", re.I)
_HEAD_NUM = re.compile(r"^\d+(\.\d+)*[.)]?\s+[A-Za-z].{2,70}$")


def is_heading_line(line: str) -> bool:
    line = line.strip()
    if not 3 <= len(line) <= 80:
        return False
    if line.endswith((".", ",", ";", ":")) and not _HEAD_KW.match(line):
        return False
    if _HEAD_KW.match(line) or _HEAD_NUM.match(line):
        return True
    if sum(c.isalpha() for c in line) < 3:
        return False
    words = line.split()
    if line.isupper() and len(words) <= 8:
        return True
    return 2 <= len(words) <= 8 and all(w[0].isupper() or not w[0].isalpha() or len(w) <= 3 for w in words)


def guess_heading(text: str) -> Optional[str]:
    lines = [l.strip() for l in text.splitlines() if l.strip()][:3]
    for l in lines:
        if is_heading_line(l):
            return l[:200]
    return None


def _ocr_available() -> bool:
    try:
        import cv2  # noqa: F401
        import numpy  # noqa: F401
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def _ocr_image_bytes(data: bytes) -> Optional[str]:
    try:
        import cv2
        import numpy as np
        import pytesseract
        img = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
        if img is None:
            return None
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
        return pytesseract.image_to_string(gray).strip() or None
    except Exception:  # missing binary, decode errors, etc.
        return None


def _describe(data: bytes, mime: str, warns: list) -> Optional[str]:
    try:
        return llm.describe_image(data, mime)
    except AIConfigError:
        warns.append("Image descriptions skipped: OPENAI_API_KEY is not set.")
    except AIServiceError as exc:
        warns.append(f"An image description failed: {exc.message}")
    return None


def _images_enabled(warns: list) -> bool:
    if not settings.enable_image_descriptions:
        return False
    if not llm.is_configured():
        warns.append("Image descriptions skipped: OPENAI_API_KEY is not set.")
        return False
    return True


# ---------- extractors ----------
def extract_pdf(path: Path, title: str, warns: list) -> list:
    try:
        import fitz
    except ImportError:
        raise ExtractionError("PyMuPDF is not installed (pip install pymupdf).")
    segs, topic, empty_pages, described = [], title, [], 0
    images_on = _images_enabled(warns)
    try:
        pdf = fitz.open(str(path))
    except Exception:
        raise ExtractionError("The PDF could not be opened (corrupt or unsupported).")
    with pdf:
        if pdf.needs_pass:
            raise ExtractionError("The PDF is password-protected.")
        ocr_ok = settings.enable_ocr and _ocr_available()
        for i, page in enumerate(pdf, start=1):
            text, ctype = page.get_text("text").strip(), "text"
            if not text and ocr_ok:
                text = _ocr_image_bytes(page.get_pixmap(dpi=200).tobytes("png")) or ""
                ctype = "ocr_text"
            if not text:
                empty_pages.append(i)
            else:
                topic = guess_heading(text) or topic
                segs.append(Segment(text, topic, page=i, content_type=ctype))
            if images_on and described < settings.max_images_per_document:
                for img in page.get_images(full=True):
                    if described >= settings.max_images_per_document:
                        break
                    try:
                        info = pdf.extract_image(img[0])
                    except Exception:
                        continue
                    if not info or min(info.get("width", 0), info.get("height", 0)) < 150:
                        continue
                    ext = str(info.get("ext", "")).lower()
                    if ext not in {"png", "jpeg", "jpg", "webp", "gif"}:
                        continue
                    desc = _describe(info["image"], f"image/{'jpeg' if ext == 'jpg' else ext}", warns)
                    if desc:
                        described += 1
                        segs.append(Segment(f"[AI-generated description of an image on page {i}] {desc}", topic,
                                            page=i, content_type="image_description"))
    if empty_pages:
        shown = ", ".join(map(str, empty_pages[:15])) + ("..." if len(empty_pages) > 15 else "")
        hint = "" if ocr_ok else " OCR is unavailable (install opencv-python-headless, pytesseract and the Tesseract binary)."
        warns.append(f"No extractable text on page(s) {shown}.{hint}")
    return segs


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def extract_pptx(path: Path, title: str, warns: list) -> list:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        raise ExtractionError("python-pptx is not installed (pip install python-pptx).")
    try:
        prs = Presentation(str(path))
    except Exception:
        raise ExtractionError("The PPTX file could not be opened (corrupt or unsupported).")
    segs, topic, described = [], title, 0
    images_on = _images_enabled(warns)
    for idx, slide in enumerate(prs.slides, start=1):
        lines, slide_title, pictures = [], None, []
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text_frame.text.strip() or None
        except Exception:
            pass
        for sh in _iter_shapes(slide.shapes):
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    lines.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures.append(sh)
        if slide_title:
            topic = slide_title[:200]
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append("Speaker notes: " + notes)
        except Exception:
            pass
        if lines:
            segs.append(Segment("\n".join(lines), topic, slide=idx, content_type="slide"))
        if images_on:
            for pic in pictures:
                if described >= settings.max_images_per_document:
                    break
                try:
                    blob, mime = pic.image.blob, pic.image.content_type
                except Exception:
                    continue
                if mime not in {"image/png", "image/jpeg", "image/gif", "image/webp"}:
                    continue
                desc = _describe(blob, mime, warns)
                if desc:
                    described += 1
                    segs.append(Segment(f"[AI-generated description of an image on slide {idx}] {desc}", topic,
                                        slide=idx, content_type="image_description"))
    return segs


def extract_text_file(path: Path, title: str, is_md: bool, warns: list) -> list:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
        warns.append("File was not valid UTF-8; decoded as Latin-1.")
    sections, topic, buf = [], title, []

    def flush():
        if buf:
            sections.append(Segment("\n\n".join(buf), topic))
            buf.clear()

    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        if is_md and lines[0].lstrip().startswith("#"):
            flush()
            topic = lines[0].lstrip("# ").strip()[:200] or topic
            rest = "\n".join(lines[1:]).strip()
            if rest:
                buf.append(rest)
        elif not is_md and len(lines) == 1 and len(block) <= 80 and not block.endswith((".", "!", "?", ",", ";", ":")):
            flush()
            topic = block[:200]
        else:
            buf.append(block)
    flush()
    return sections


def extract_video(path: Path, title: str, warns: list) -> list:
    groups = video_service.transcribe_video(path, warns)
    return [Segment(g["text"], title, start_sec=g["start_sec"], end_sec=g["end_sec"],
                    content_type="transcript", presplit=True) for g in groups]


def extract(path: Path, doc: Document, warns: list) -> list:
    ext, title = "." + doc.file_type, Path(doc.filename).stem[:200]
    if ext == ".pdf":
        return extract_pdf(path, title, warns)
    if ext == ".pptx":
        return extract_pptx(path, title, warns)
    if ext in {".txt", ".md"}:
        return extract_text_file(path, title, ext == ".md", warns)
    if ext in VIDEO_EXT:
        return extract_video(path, title, warns)
    raise ExtractionError(f"Unsupported file type '{ext}'.")


# ---------- upload ----------
def _check_signature(ext: str, head: bytes) -> None:
    if ext == ".pdf" and b"%PDF-" not in head[:1024]:
        raise BadRequestError("File content is not a valid PDF.")
    if ext == ".pptx" and head[:2] != b"PK":
        raise BadRequestError("File content is not a valid PPTX.")
    if ext in {".txt", ".md"} and b"\x00" in head:
        raise BadRequestError("File looks binary, not text.")


def create_document_from_upload(db: Session, user_id: str, upload) -> Document:
    display = safe_filename(upload.filename or "")
    ext = Path(display).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise BadRequestError(f"Unsupported file type '{ext or 'none'}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    root = settings.upload_dir.resolve()
    root.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}{ext}"
    dest = (root / stored).resolve()
    if root not in dest.parents:
        raise BadRequestError("Invalid upload path.")
    size, head = 0, b""
    try:
        with open(dest, "wb") as out:
            while True:
                block = upload.file.read(1024 * 1024)
                if not block:
                    break
                if size == 0:
                    head = block[:4096]
                size += len(block)
                if size > settings.max_upload_bytes:
                    raise PayloadTooLargeError(f"File exceeds the {settings.max_upload_size_mb} MB limit.")
                out.write(block)
        if size == 0:
            raise BadRequestError("The uploaded file is empty.")
        _check_signature(ext, head)
    except Exception:
        dest.unlink(missing_ok=True)
        raise
    doc = Document(id=new_id(), user_id=user_id, filename=display, stored_name=stored, file_type=ext.lstrip("."),
                   size_bytes=size, status="processing", warnings=[])
    db.add(doc)
    db.add(ActivityLog(user_id=user_id, kind="upload", detail={"filename": display}))
    db.commit()
    db.refresh(doc)
    return doc


# ---------- processing ----------
def index_document(db: Session, doc: Document) -> None:
    rows = db.scalars(select(Chunk).where(Chunk.document_id == doc.id).order_by(Chunk.chunk_index)).all()
    items = [(c.id, c.text, embedding_service.build_metadata(c, doc.filename)) for c in rows]
    embedding_service.index_chunks(items)
    for c in rows:
        c.indexed = True


def _clear_chunks(db: Session, doc: Document) -> None:
    try:
        embedding_service.delete_document(doc.id)
    except Exception:
        log.exception("Could not clear vectors for %s", doc.id)
    for c in db.scalars(select(Chunk).where(Chunk.document_id == doc.id)).all():
        db.delete(c)
    doc.chunk_count = 0
    db.commit()


def process_document(document_id: str) -> None:
    """Runs in a background task with its own session. Never raises."""
    db = SessionLocal()
    try:
        doc = db.get(Document, document_id)
        if doc is None:
            return
        doc.status, doc.error = "processing", None
        db.commit()
        _clear_chunks(db, doc)
        warns: list = []
        try:
            segments = extract(settings.upload_dir / doc.stored_name, doc, warns)
        except (ExtractionError, AIConfigError) as exc:
            doc.status, doc.error, doc.warnings = "failed", exc.message, list(dict.fromkeys(warns))
            db.commit()
            return
        except Exception:
            log.exception("Extraction crashed for %s", document_id)
            doc.status, doc.error = "failed", "Unexpected error while extracting content (see server logs)."
            doc.warnings = list(dict.fromkeys(warns))
            db.commit()
            return
        idx, size, overlap = 0, settings.chunk_size, settings.effective_overlap
        for seg in segments:
            pieces = [seg.text.strip()] if seg.presplit else chunk_text(seg.text, size, overlap)
            for p in pieces:
                if not p:
                    continue
                db.add(Chunk(id=f"{doc.id}-{idx:05d}", document_id=doc.id, user_id=doc.user_id, chunk_index=idx,
                             text=p, page=seg.page, slide=seg.slide, start_sec=seg.start_sec, end_sec=seg.end_sec,
                             topic=(seg.topic or "General")[:200], content_type=seg.content_type, indexed=False))
                idx += 1
        doc.chunk_count, doc.warnings = idx, list(dict.fromkeys(warns))
        if idx == 0:
            doc.status, doc.error = "failed", "No extractable text was found in this file."
            db.commit()
            return
        db.commit()
        if not llm.is_configured():
            doc.status = "awaiting_index"
            doc.error = "Text was extracted, but OPENAI_API_KEY is not set so it cannot be indexed yet. Set it and call POST /api/documents/{id}/reprocess."
            db.commit()
            return
        try:
            index_document(db, doc)
            doc.status = "ready"
        except AIServiceError as exc:
            doc.status, doc.error = "awaiting_index", f"Indexing failed: {exc.message} Retry with POST /api/documents/{{id}}/reprocess."
        except Exception:
            log.exception("Indexing crashed for %s", document_id)
            doc.status, doc.error = "awaiting_index", "Indexing failed unexpectedly (see server logs)."
        db.commit()
    finally:
        db.close()


def delete_document(db: Session, doc: Document) -> None:
    try:
        embedding_service.delete_document(doc.id)
    except Exception:
        log.exception("Vector deletion failed for %s (orphans are ignored at retrieval time)", doc.id)
    (settings.upload_dir / doc.stored_name).unlink(missing_ok=True)
    db.delete(doc)  # cascades to chunks
    db.commit()